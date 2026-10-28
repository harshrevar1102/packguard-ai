"""Content-sheet extraction.

Nothing here assumes a fixed worksheet or cell address. Fields are located by their *labels*
(a cell whose text matches a known label, with the value in the cell(s) to its right), and the
nutrition table is located by its header cell. Every extracted value keeps its worksheet + cell.
Formula cells are never used as reference values (they are derived, not authored content)."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Union

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.models import (
    NutrientReference,
    ReferenceExtraction,
    ReferenceField,
    ReferenceStatus,
    SourceRef,
    UnmappedRow,
)
from app.services.text_utils import collapse, slug

class ProcessingError(Exception):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


# key -> label regexes (matched against the lower-cased, colon-stripped label text)
FIELD_DEFS: dict[str, list[str]] = {
    "label_identifier": [r"^label identifier$"],
    "sheet_status": [r"^status$"],
    "product_name": [r"^product name$", r"^product$", r"^item name$"],
    "brand": [r"^brand( name)?$"],
    "bpn": [r"^bpn$"],
    "upc": [r"^upc\b", r"^barcode$", r"^gtin\b"],
    "legal_name": [r"^legal name\b", r"^statement of identity$"],
    "allergens": [r"^allergen information$", r"^allergens?$", r"^allergen declaration$"],
    "storage": [r"^storage instructions?$", r"^storage$"],
    "ingredients": [r"^ingredients?( list)?$"],
    "marketing_copy": [r"^marketing copy$"],
    "dietary": [r"^dietary information$", r"^dietary claims?$"],
    "claims": [r"claims (&|and) certifications", r"^claims?$"],
    "servings_per_container": [r"^number of servings per (package|container)$", r"^servings per (package|container)$", r"^servings$"],
    "net_weight": [r"\bnet (weight|contents?)\b", r"^package size\b", r"^net wt\b"],
    "serving_size": [r"^serving size$"],
    "nutrition_footnote": [r"^nutrition panel bottom statement$", r"^footnote$"],
}
_COMPILED = {k: [re.compile(p) for p in v] for k, v in FIELD_DEFS.items()}

# Common nutrient labels when not formatted in a standard grid table
_NUTRIENT_LABELS = {
    "calories": "Calories",
    "energy": "Calories",
    "total fat": "Total Fat",
    "saturated fat": "Saturated Fat",
    "trans fat": "Trans Fat",
    "cholesterol": "Cholesterol",
    "sodium": "Sodium",
    "total carbohydrate": "Total Carbohydrate",
    "carbohydrates": "Total Carbohydrate",
    "dietary fiber": "Dietary Fiber",
    "fiber": "Dietary Fiber",
    "total sugars": "Total Sugars",
    "sugars": "Total Sugars",
    "added sugars": "Includes Added Sugars",
    "protein": "Protein",
    "vitamin d": "Vitamin D",
    "calcium": "Calcium",
    "iron": "Iron",
    "potassium": "Potassium",
}

# Wording that means "this value is not settled" (applied to the cell value and its note).
REVIEW_RE = re.compile(
    r"\b(verify|confirm|assumed|assumption|constructed|reconstructed|unconfirmed|to be confirmed|"
    r"tbc|tbd|please check|not clearly|illegible|placeholder|needs? review|draft)\b",
    re.I,
)
_DRAFT_RE = re.compile(r"draft|in review|unapproved|pending|work in progress|\bwip\b", re.I)
_APPROVED_RE = re.compile(r"\bapproved\b|\breleased\b|\bfinal\b", re.I)
_EMPTY_VALUES = {"", "n/a", "na", "-", "—", "none"}


def _label_key(text: str) -> str:
    return collapse(text.lower()).rstrip(":").strip()


def _match_field(label: str) -> str | None:
    lab = _label_key(label)
    for key, pats in _COMPILED.items():
        if any(p.search(lab) for p in pats):
            return key
    return None


def _text(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


@dataclass
class _Sheet:
    name: str
    cells: dict  # (row, col) -> str  (non-formula, non-empty)
    formulas: set  # (row, col)


def _read_sheets(data_or_path: Union[bytes, Path, str]) -> list[_Sheet]:
    try:
        if isinstance(data_or_path, (Path, str)):
            wb = load_workbook(str(data_or_path), data_only=False)
        else:
            wb = load_workbook(io.BytesIO(data_or_path), data_only=False)
    except Exception as exc:
        raise ProcessingError(
            "The content sheet could not be read as an Excel (.xlsx) workbook. "
            "The file may be corrupted or not a real .xlsx file."
        ) from exc
    sheets: list[_Sheet] = []
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible":
            continue
        cells, formulas = {}, set()
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                t = _text(c.value)
                if not t:
                    continue
                if isinstance(c.value, str) and t.startswith("="):
                    formulas.add((c.row, c.column))
                else:
                    cells[(c.row, c.column)] = t
        sheets.append(_Sheet(ws.title, cells, formulas))
    if not sheets or not any(s.cells for s in sheets):
        raise ProcessingError("The workbook contains no readable, populated worksheets.")
    return sheets


def _ref(sheet: str, r: int, c: int) -> SourceRef:
    return SourceRef(sheet=sheet, cell=f"{get_column_letter(c)}{r}")


def extract_reference(data_or_path: Union[bytes, Path, str]) -> ReferenceExtraction:
    sheets = _read_sheets(data_or_path)
    fields: dict[str, ReferenceField] = {}
    warnings: list[str] = []
    consumed: dict[str, set] = {s.name: set() for s in sheets}

    for sh in sheets:
        for (r, c), text in sorted(sh.cells.items()):
            key = _match_field(text)
            if not key:
                continue
            # value = first populated cell within two columns to the right on the same row
            vcol = next((c + d for d in (1, 2) if (r, c + d) in sh.cells), None)
            if vcol is None:
                continue
            if key in fields:
                warnings.append(
                    f"Label “{text}” appears more than once; the first occurrence "
                    f"({fields[key].label_source.display()}) was used and {sh.name}!{get_column_letter(c)}{r} was ignored."
                )
                continue
            value = sh.cells[(r, vcol)]
            note = note_src = None
            ncell = (r, vcol + 1)
            if ncell in sh.cells and _match_field(sh.cells[ncell]) is None:
                note = sh.cells[ncell]
                note_src = _ref(sh.name, *ncell)
                consumed[sh.name].add(ncell)
            reason = None
            if note and REVIEW_RE.search(note):
                reason = note
            elif REVIEW_RE.search(value) and len(value) < 200:
                reason = f"Value itself indicates uncertainty: “{value}”"
            fields[key] = ReferenceField(
                key=key,
                label=collapse(text),
                value=value,
                source=_ref(sh.name, r, vcol),
                label_source=_ref(sh.name, r, c),
                note=note,
                note_source=note_src,
                flagged=reason is not None,
                flag_reason=reason,
            )
            consumed[sh.name].update({(r, c), (r, vcol)})

    # ---- sheet-level approval status
    st = fields.get("sheet_status")
    if st:
        approved = False if _DRAFT_RE.search(st.value) else (True if _APPROVED_RE.search(st.value) else None)
        status = ReferenceStatus(text=st.value, source=st.source, note=st.note, approved=approved)
    else:
        status = ReferenceStatus(text=None, approved=None)
        warnings.append("No explicit Status field was found in the workbook; approval status is unknown.")

    # ---- nutrition table
    nutrition, title, nwarn = _extract_nutrition(sheets, consumed)
    warnings.extend(nwarn)

    # ---- fallback: if no nutrition table was found, check for key-value nutrient rows
    if not nutrition:
        nutrition = _extract_standalone_nutrients(sheets, consumed)

    if not (set(fields) - {"sheet_status"}) and not nutrition:
        raise ProcessingError(
            "No supported fields (e.g. Ingredients, Allergen Information, Net Weight, Nutrition table) "
            "could be identified in this workbook. Check that the sheet uses labelled fields."
        )

    unmapped: list[UnmappedRow] = []
    for sh in sheets:
        rows: dict[int, list] = {}
        for (r, c), t in sh.cells.items():
            rows.setdefault(r, []).append((c, t))
        for r in sorted(rows):
            cs = sorted(rows[r])
            un = [(c, t) for c, t in cs if (r, c) not in consumed[sh.name]]
            if len(cs) >= 2 and un:
                txt = " | ".join(f"{get_column_letter(c)}{r}: {collapse(t)[:140]}" for c, t in un)
                unmapped.append(UnmappedRow(sheet=sh.name, row=r, text=txt))
            if len(unmapped) >= 60:
                break

    return ReferenceExtraction(
        sheets=[s.name for s in sheets],
        status=status,
        fields=fields,
        nutrition=nutrition,
        nutrition_title=title,
        warnings=warnings,
        unmapped_rows=unmapped,
    )


def _extract_nutrition(sheets: list[_Sheet], consumed: dict[str, set]):
    warnings: list[str] = []
    for sh in sheets:
        header = next(
            ((r, c) for (r, c), t in sorted(sh.cells.items()) if _label_key(t) in ("nutrient composition", "nutrition facts", "nutrition information")), None
        )
        if not header:
            continue
        hr, hc = header
        # section title: nearest populated cell above header in the same column
        title = next((sh.cells[(r, hc)] for r in range(hr - 1, max(0, hr - 4), -1) if (r, hc) in sh.cells), None)
        out: list[NutrientReference] = []
        blank = 0
        n_formula = 0
        r = hr + 1
        while blank < 2 and r < hr + 80:
            if (r, hc) not in sh.cells:
                blank += 1
                r += 1
                continue
            blank = 0
            label = sh.cells[(r, hc)]
            cols = [c for c in range(hc + 1, hc + 3) if (r, c) in sh.cells]
            amount = sh.cells.get((r, hc + 1))
            dv = sh.cells.get((r, hc + 2))
            n_formula += sum(1 for (fr, _fc) in sh.formulas if fr == r)
            out.append(
                NutrientReference(
                    key=slug(label),
                    label=collapse(label),
                    amount=amount,
                    amount_source=_ref(sh.name, r, hc + 1) if amount is not None else None,
                    daily_value=dv,
                    daily_value_source=_ref(sh.name, r, hc + 2) if dv is not None else None,
                    label_source=_ref(sh.name, r, hc),
                )
            )
            consumed[sh.name].update({(r, hc)} | {(r, c) for c in cols})
            r += 1
        consumed[sh.name].add((hr, hc))
        if title and "as printed" not in title.lower():
            warnings.append(
                f"The nutrition table title (“{title}”) does not say the values are as printed on the pack; "
                "treat nutrition comparisons with caution."
            )
        if n_formula:
            warnings.append(
                f"{n_formula} formula cell(s) in the nutrition table (e.g. scaled to other pack sizes) were not used; "
                f"only authored values in {get_column_letter(hc + 1)}–{get_column_letter(hc + 2)} are compared."
            )
        return out, title, warnings
    return [], None, warnings


def _extract_standalone_nutrients(sheets: list[_Sheet], consumed: dict[str, set]) -> list[NutrientReference]:
    """Fallback when nutrients are declared in key-value rows rather than under a table header."""
    out: list[NutrientReference] = []
    seen = set()
    for sh in sheets:
        for (r, c), text in sorted(sh.cells.items()):
            if (r, c) in consumed[sh.name]:
                continue
            norm_lbl = _label_key(text)
            matched_nutrient = _NUTRIENT_LABELS.get(norm_lbl)
            if not matched_nutrient:
                continue
            k = slug(matched_nutrient)
            if k in seen:
                continue
            vcol = next((c + d for d in (1, 2) if (r, c + d) in sh.cells), None)
            if vcol is None:
                continue
            amount = sh.cells[(r, vcol)]
            dv_col = vcol + 1
            dv = sh.cells.get((r, dv_col)) if (r, dv_col) in sh.cells and "%" in sh.cells.get((r, dv_col), "") else None
            out.append(
                NutrientReference(
                    key=k,
                    label=matched_nutrient,
                    amount=amount,
                    amount_source=_ref(sh.name, r, vcol),
                    daily_value=dv,
                    daily_value_source=_ref(sh.name, r, dv_col) if dv else None,
                    label_source=_ref(sh.name, r, c),
                )
            )
            consumed[sh.name].update({(r, c), (r, vcol)})
            if dv:
                consumed[sh.name].add((r, dv_col))
            seen.add(k)
    return out


def is_empty_value(v: str | None) -> bool:
    return v is None or v.strip().lower() in _EMPTY_VALUES
